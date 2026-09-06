from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
import json
import os
import zipfile

import fitz
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as PptxInches


ROOT = Path(__file__).resolve().parent
FIXTURES = ROOT / "fixtures"
FIXED_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)
ZIP_TIME = (2026, 1, 1, 0, 0, 0)


def _image(path: Path, label: str, *, color: tuple[int, int, int] = (238, 247, 255)) -> None:
    image = Image.new("RGB", (1200, 800), color)
    draw = ImageDraw.Draw(image)
    draw.rectangle((80, 80, 1120, 720), outline=(28, 89, 140), width=8)
    draw.ellipse((180, 220, 520, 560), outline=(24, 122, 90), width=12)
    draw.line((520, 390, 920, 390), fill=(170, 70, 40), width=12)
    draw.text((120, 120), label, fill=(10, 30, 50), stroke_width=1)
    draw.text((120, 650), "Synthetic fixture - no personal data", fill=(60, 60, 60))
    image.save(path)


def _native_pdf(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 80), "Synthetic Biology Course", fontsize=20)
    page.insert_text((72, 125), "Chapter 1  Cell structure", fontsize=16)
    page.insert_textbox(
        fitz.Rect(72, 150, 520, 360),
        "Cells are the basic structural and functional units of life.\n"
        "The plasma membrane regulates transport and communication.\n"
        "This file contains a native selectable text layer.",
        fontsize=12,
    )
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 80), "Chapter 2  Genetics", fontsize=18)
    page.insert_textbox(
        fitz.Rect(72, 120, 520, 320),
        "DNA stores hereditary information. During transcription, DNA is used as a template for RNA.",
        fontsize=12,
    )
    document.set_metadata({"title": "Stage 0 native text fixture", "author": "CloudPath synthetic fixture"})
    document.save(path, no_new_id=True)
    document.close()


def _complex_pdf(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((50, 55), "Two-column lesson with table and formula", fontsize=17)
    page.insert_textbox(
        fitz.Rect(50, 90, 280, 390),
        "LEFT COLUMN\nMitosis produces genetically similar daughter cells.\n"
        "Prophase is followed by metaphase, anaphase, and telophase.",
        fontsize=11,
    )
    page.insert_textbox(
        fitz.Rect(315, 90, 545, 390),
        "RIGHT COLUMN\nMeiosis contains two divisions and produces haploid cells.\n"
        "Independent assortment increases variation.",
        fontsize=11,
    )
    page.insert_text((50, 440), "Formula: P(A and B) = P(A) * P(B)", fontsize=12)
    rows = [
        ("Stage", "Chromosome state"),
        ("G1", "unreplicated"),
        ("S", "DNA replication"),
        ("G2", "replicated"),
    ]
    x_positions = (50, 250, 520)
    y = 490
    for row_index, row in enumerate(rows):
        page.draw_line((x_positions[0], y), (x_positions[-1], y))
        page.insert_text((x_positions[0] + 5, y + 20), row[0], fontsize=10)
        page.insert_text((x_positions[1] + 5, y + 20), row[1], fontsize=10)
        y += 32
        if row_index == len(rows) - 1:
            page.draw_line((x_positions[0], y), (x_positions[-1], y))
    for x in x_positions:
        page.draw_line((x, 490), (x, y))
    document.set_metadata({"title": "Stage 0 complex layout fixture", "author": "CloudPath synthetic fixture"})
    document.save(path, no_new_id=True)
    document.close()


def _scanned_pdf(path: Path, image_path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_image(fitz.Rect(40, 100, 555, 443), filename=str(image_path))
    document.set_metadata({"title": "Stage 0 scanned PDF fixture"})
    document.save(path, no_new_id=True)
    document.close()


def _mixed_pdf(path: Path, image_path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 90), "Native text page: Photosynthesis converts light energy.", fontsize=13)
    page = document.new_page(width=595, height=842)
    page.insert_image(fitz.Rect(40, 100, 555, 443), filename=str(image_path))
    document.set_metadata({"title": "Stage 0 mixed PDF fixture"})
    document.save(path, no_new_id=True)
    document.close()


def _encrypted_pdf(path: Path) -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 90), "Encrypted synthetic fixture", fontsize=14)
    document.save(
        path,
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="stage0-owner",
        user_pw="stage0-user",
        permissions=0,
        no_new_id=True,
    )
    document.close()


def _normalize_ooxml(path: Path) -> None:
    replacement = path.with_suffix(path.suffix + ".normalized")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(replacement, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    os.replace(replacement, path)


def _docx(path: Path, image_path: Path) -> None:
    document = Document()
    document.core_properties.title = "Synthetic biology handout"
    document.core_properties.author = "CloudPath synthetic fixture"
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.add_heading("Synthetic Biology Handout", level=0)
    document.add_heading("1. Cell membrane", level=1)
    document.add_paragraph("The cell membrane controls exchange between the cell and its environment.")
    document.add_heading("1.1 Transport comparison", level=2)
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Process", "Energy"), ("Diffusion", "No ATP"), ("Active transport", "Uses ATP")], strict=True):
        row.cells[0].text, row.cells[1].text = values
    document.add_picture(str(image_path), width=Inches(3.5))
    document.add_paragraph("Figure 1. Synthetic cell diagram.")
    document.save(path)
    _normalize_ooxml(path)


def _pptx(path: Path, image_path: Path) -> None:
    presentation = Presentation()
    presentation.core_properties.title = "Synthetic biology slides"
    presentation.core_properties.author = "CloudPath synthetic fixture"
    presentation.core_properties.created = FIXED_TIME
    presentation.core_properties.modified = FIXED_TIME
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Cellular respiration"
    slide.placeholders[1].text = "Synthetic stage 0 fixture"
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "ATP production"
    textbox = slide.shapes.add_textbox(PptxInches(0.7), PptxInches(1.4), PptxInches(5.0), PptxInches(2.0))
    textbox.text_frame.text = "Glycolysis occurs in the cytoplasm.\nThe citric acid cycle occurs in mitochondria."
    slide.shapes.add_picture(str(image_path), PptxInches(6.0), PptxInches(1.5), width=PptxInches(3.0))
    presentation.save(path)
    _normalize_ooxml(path)


def _xlsx(path: Path) -> None:
    workbook = Workbook()
    workbook.properties.title = "Synthetic experiment workbook"
    workbook.properties.creator = "CloudPath synthetic fixture"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    sheet = workbook.active
    sheet.title = "Growth Data"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Synthetic Plant Growth"
    sheet["A1"].font = Font(bold=True)
    sheet.append(["Day", "Control", "Treatment"])
    for values in [(1, 2.0, 2.1), (2, 2.8, 3.2), (3, 3.5, 4.4)]:
        sheet.append(values)
    sheet["D2"] = "Difference"
    sheet["D3"] = "=C3-B3"
    notes = workbook.create_sheet("Notes")
    notes.append(["Observation", "Leaves remained green in both groups."])
    workbook.save(path)
    _normalize_ooxml(path)


def _security_zip(path: Path, *, traversal: bool = False, high_ratio: bool = False) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types></Types>")
        archive.writestr("word/document.xml", "<document>synthetic</document>")
        if traversal:
            archive.writestr("../escape.txt", "must never be extracted")
        if high_ratio:
            archive.writestr("word/high_ratio.bin", b"A" * (2 * 1024 * 1024))


def _write_manifest(entries: list[dict[str, object]]) -> None:
    for entry in entries:
        file_path = FIXTURES / str(entry["file"])
        entry["bytes"] = file_path.stat().st_size
        entry["sha256"] = sha256(file_path.read_bytes()).hexdigest()
    payload = {
        "schema_version": 1,
        "generated_at": "2026-07-13",
        "contains_personal_data": False,
        "entries": entries,
    }
    (ROOT / "fixture_manifest.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    for existing in FIXTURES.iterdir():
        if existing.is_file():
            existing.unlink()

    source_image = FIXTURES / "source_diagram.png"
    _image(source_image, "CELL DIAGRAM")
    _native_pdf(FIXTURES / "native_text.pdf")
    _complex_pdf(FIXTURES / "multicolumn_table_formula.pdf")
    _scanned_pdf(FIXTURES / "scanned.pdf", source_image)
    _mixed_pdf(FIXTURES / "mixed.pdf", source_image)
    _encrypted_pdf(FIXTURES / "encrypted.pdf")

    base = Image.open(source_image)
    base.save(FIXTURES / "sample.jpg", quality=92)
    base.save(FIXTURES / "sample.jpeg", quality=92)
    base.save(FIXTURES / "sample.webp", quality=90)
    base.save(FIXTURES / "sample.bmp")
    base.save(FIXTURES / "sample.jp2")
    frame_2 = Image.new("RGB", base.size, (255, 240, 220))
    ImageDraw.Draw(frame_2).text((120, 120), "ANIMATED FRAME 2", fill=(80, 20, 10))
    base.save(FIXTURES / "animated.gif", save_all=True, append_images=[frame_2], duration=250, loop=0)
    base.save(FIXTURES / "multipage.tiff", save_all=True, append_images=[frame_2])
    base.save(FIXTURES / "sample.tif")

    _docx(FIXTURES / "sample.docx", source_image)
    _pptx(FIXTURES / "sample.pptx", source_image)
    _xlsx(FIXTURES / "sample.xlsx")
    _security_zip(FIXTURES / "path_traversal.docx", traversal=True)
    _security_zip(FIXTURES / "high_compression_ratio.docx", high_ratio=True)
    (FIXTURES / "spoofed.docx").write_text("This is plain text, not an OOXML package.\n", encoding="utf-8")
    (FIXTURES / "corrupt.pdf").write_bytes(b"%PDF-1.7\ncorrupt synthetic fixture\n")
    (FIXTURES / "corrupt.png").write_bytes(b"\x89PNG\r\n\x1a\ncorrupt")

    _write_manifest(
        [
            {"file": "native_text.pdf", "category": "valid", "scenario": "native text PDF"},
            {"file": "multicolumn_table_formula.pdf", "category": "valid", "scenario": "multi-column, table, formula"},
            {"file": "scanned.pdf", "category": "valid", "scenario": "image-only PDF"},
            {"file": "mixed.pdf", "category": "valid", "scenario": "native and scanned pages"},
            {"file": "encrypted.pdf", "category": "reject", "scenario": "encrypted PDF"},
            {"file": "source_diagram.png", "category": "valid", "scenario": "PNG"},
            {"file": "sample.jpg", "category": "valid", "scenario": "JPG"},
            {"file": "sample.jpeg", "category": "valid", "scenario": "JPEG"},
            {"file": "sample.webp", "category": "valid", "scenario": "WEBP"},
            {"file": "sample.bmp", "category": "valid", "scenario": "BMP"},
            {"file": "sample.jp2", "category": "valid", "scenario": "JPEG 2000"},
            {"file": "animated.gif", "category": "decision", "scenario": "two-frame GIF"},
            {"file": "multipage.tiff", "category": "decision", "scenario": "two-page TIFF"},
            {"file": "sample.tif", "category": "decision", "scenario": "TIF alias"},
            {"file": "sample.docx", "category": "valid", "scenario": "DOCX headings, table, image"},
            {"file": "sample.pptx", "category": "valid", "scenario": "PPTX slides, image"},
            {"file": "sample.xlsx", "category": "valid", "scenario": "XLSX sheets, merged cells, formula"},
            {"file": "path_traversal.docx", "category": "reject", "scenario": "OOXML path traversal"},
            {"file": "high_compression_ratio.docx", "category": "reject", "scenario": "high compression ratio"},
            {"file": "spoofed.docx", "category": "reject", "scenario": "extension spoofing"},
            {"file": "corrupt.pdf", "category": "reject", "scenario": "corrupt PDF"},
            {"file": "corrupt.png", "category": "reject", "scenario": "corrupt PNG"},
        ]
    )


if __name__ == "__main__":
    main()
